"""Generatore del file xlsx di prova (3 fogli, 20 colonne, intestazione alla riga 2)."""


def ean13(prefix12: str) -> str:
    total = sum(int(c) * (3 if i % 2 else 1) for i, c in enumerate(prefix12))
    return prefix12 + str((10 - total % 10) % 10)


def build(path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Titolo"
    ws.append(["Note"])
    ws.append(["solo testo"])
    dati = wb.create_sheet("Dati")
    dati["A1"] = "Listino 2026 (riga titolo, non e' l'intestazione)"
    header = ["Codice", "EAN", "Descrizione", "Layout", "Misto", None, "Descrizione", None, "Prezzo"] + \
             [f"Extra {i}" for i in range(1, 12)]
    assert len(header) == 20
    dati.append([None])  # riga 2 verra' sostituita sotto
    for c, h in enumerate(header, start=1):
        dati.cell(row=2, column=c, value=h)
    for i in range(1, 61):
        mixed = ean13(f"59012341{i:04d}") if i % 3 else f"AB-{i}x"
        row = [f"COD{i:04d}", ean13(f"59012340{i:04d}"), f"Articolo di prova numero {i}", f"A{i % 9}B",
               mixed, "x", f"Seconda descrizione {i}", None, 10.5 + i] + [f"e{i}-{k}" for k in range(11)]
        for c, v in enumerate(row, start=1):
            dati.cell(row=2 + i, column=c, value=v)
    wb.create_sheet("Vuoto")
    wb.save(path)
    return str(path)
