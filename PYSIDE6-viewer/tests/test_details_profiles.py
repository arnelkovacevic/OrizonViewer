import os

from app import file_profiles, xlsx_reader
from app.details import DetailSpec, count_invalid, is_valid, suggest_type
from app.file_profiles import Profile, ProfileStore, resolve, signature


def test_validation_rules():
    assert is_valid("code128", "90702010185217i")
    assert is_valid("ean13", "5901234123457") and not is_valid("ean13", "5901234123458")
    assert is_valid("ean13", "590123412345")            # 12 cifre: checksum calcolato
    assert not is_valid("ean13", "ABC")
    assert is_valid("ean8", "96385074") and not is_valid("ean8", "96385075")
    assert is_valid("upca", "036000291452")
    assert is_valid("code39", "abc-12") and not is_valid("code39", "abc_12")
    assert is_valid("itf", "12345") and not is_valid("itf", "12a")
    assert is_valid("qr", "qualsiasi cosa ✓")
    assert not is_valid("code128", "")


def test_suggestion_and_warning():
    ean = ["5901234123457", "4006381333931"]
    assert suggest_type(ean) == "ean13"
    assert suggest_type(ean + ["AB-1"]) == "code128"
    assert suggest_type([]) == "code128"
    assert count_invalid("ean13", ean + ["AB-1", ""]) == (1, 3)


def test_reader_names_and_header_row(fixture_xlsx):
    sheets = xlsx_reader.list_sheets(fixture_xlsx)
    assert [s.name for s in sheets] == ["Titolo", "Dati", "Vuoto"]
    pv = xlsx_reader.read_preview(fixture_xlsx, "Dati", header_row=2)
    assert len(pv.names) == 20
    assert pv.names[5] == "Colonna 6" and pv.names[7] == "Colonna 8"      # intestazioni vuote
    assert pv.names[2] == "Descrizione" and pv.names[6] == "Descrizione (2)"  # duplicati
    assert pv.samples[1].isdigit() and len(pv.samples[1]) == 13           # EAN come testo, senza «.0»
    assert len(pv.rows) == 60


def test_load_only_needed_columns(fixture_xlsx):
    ds = xlsx_reader.load_dataset(fixture_xlsx, "Dati", 2, ["Descrizione (2)", "EAN"])
    assert ds.columns == ["Descrizione (2)", "EAN"] and len(ds) == 60
    assert ds.rows[0][0] == "Seconda descrizione 1"
    assert ds.search("seconda%descrizione%7") != []


def test_profile_roundtrip_and_match(tmp_path, fixture_xlsx):
    store = ProfileStore(str(tmp_path / "profiles.json"))
    pv = xlsx_reader.read_preview(fixture_xlsx, "Dati", 2)
    spec = DetailSpec(enabled=True, code_col="EAN", barcode_type="ean13", above_col="Descrizione",
                      hl_col="Layout", hl_color="blu")
    prof = Profile(fixture_xlsx, "Dati", 2, pv.names, ["Codice", "EAN", "Layout"], spec, {"Codice": 120.0})
    assert prof.needed_columns() == ["Codice", "EAN", "Layout", "Descrizione"]
    store.upsert(prof)

    store2 = ProfileStore(str(tmp_path / "profiles.json"))      # riletto da disco
    assert store2.count() == 1
    got = store2.all()[0]
    assert got.columns == prof.columns and got.details.hl_color == "blu" and got.widths == {"Codice": 120.0}
    assert resolve(fixture_xlsx, store2).kind == "applied"

    # file copiato altrove con la stessa struttura: il profilo si applica
    import shutil
    other = str(tmp_path / "copia.xlsx")
    shutil.copy(fixture_xlsx, other)
    m = resolve(other, store2)
    assert m.kind == "applied" and m.profile.path == other


def test_changed_structure_and_none(tmp_path, fixture_xlsx):
    from openpyxl import load_workbook
    store = ProfileStore(str(tmp_path / "p.json"))
    assert resolve(fixture_xlsx, store).kind == "none"
    pv = xlsx_reader.read_preview(fixture_xlsx, "Dati", 2)
    store.upsert(Profile(fixture_xlsx, "Dati", 2, pv.names, ["Codice", "Layout"]))
    wb = load_workbook(fixture_xlsx)
    wb["Dati"].cell(row=2, column=4, value="Posizione")      # «Layout» rinominata
    wb.save(fixture_xlsx)
    m = resolve(fixture_xlsx, store)
    assert m.kind == "changed" and m.profile.columns == ["Codice", "Layout"]


def test_signature_depends_on_sheet_and_columns():
    assert signature("A", ["x", "y"]) != signature("B", ["x", "y"])
    assert signature("A", ["x", "y"]) != signature("A", ["y", "x"])
    assert signature("A", ["x", "y"]) == signature("A", ["x", "y"])
