"""Test GUI offscreen: wizard, profili salvati, popup. I modali si pilotano con run_modal."""
import os

import pytest
from PySide6.QtCore import Qt

from app import config
from app.detail_dialog import DetailDialog, NoDetailsDialog
from app.file_profiles import ProfileStore
from app.pin_dialog import PinDialog
from app.settings_dialog import TAB_FILE, SettingsDialog
from app.wizard.wizard_dialog import WizardDialog
from gui_helpers import FakeScreen, pump, run_modal, wait_loaded

MM16 = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), "mm16.xlsx")


@pytest.fixture()
def win(qapp, tmp_path, monkeypatch):
    import app.main_window as mw
    monkeypatch.setattr(config, "ORG_NAME", "Orizon-AIX-PYTEST")
    monkeypatch.setattr(mw, "ProfileStore", lambda: ProfileStore(str(tmp_path / "profiles.json")))
    monkeypatch.setattr(PinDialog, "ask", staticmethod(lambda parent, prompt="": True))
    monkeypatch.setattr(mw.MainWindow, "screen", lambda self: FakeScreen(1024, 600))
    w = mw.MainWindow(1.0)
    w.settings._s.clear()
    w.resize(1024, 600)
    w.show()
    yield w
    w.settings._s.clear()
    w._allow_close = True
    w.close()


def _wizard_on_fixture(d):
    d.step_sheet.list.setCurrentRow(1)              # «Dati»
    d.step_sheet.stepper.plus.click()               # riga intestazione 2
    d.btn_next.click()
    assert d.step == 1
    names = d.step_cols.all_in_order()
    assert len(names) == 20 and "Colonna 6" in names and "Descrizione (2)" in names
    d.step_cols.set_all(False)
    assert not d.btn_next.isEnabled()               # serve almeno 1 colonna
    for r in d.step_cols._rows:
        if r.name in ("Codice", "EAN", "Layout", "Misto"):
            r.check.setChecked(True)
    [r for r in d.step_cols._rows if r.name == "Layout"][0].up.click()
    for _ in range(10):
        [r for r in d.step_cols._rows if r.name == "Layout"][0].up.click()
    assert d.step_cols.selected_in_order() == ["Layout", "Codice", "EAN", "Misto"]
    d.btn_next.click()
    sd = d.step_details
    sd.toggle.setChecked(True)
    sd.code_col.setCurrentIndex(sd.code_col.findData("EAN"))
    sd._on_code_col_user(0)
    assert sd.btype.currentData() == "ean13" and sd.warn.text() == ""
    sd.code_col.setCurrentIndex(sd.code_col.findData("Misto"))
    sd._on_code_col_user(0)
    assert sd.btype.currentData() == "code128"
    sd.btype.setCurrentIndex(sd.btype.findData("ean13"))
    assert sd.warn.text() == "20 valori su 60 non sono validi per EAN-13"
    sd.hl_col.setCurrentIndex(sd.hl_col.findData("Layout"))
    sd.above_col.setCurrentIndex(sd.above_col.findData("Descrizione"))
    d.btn_next.click()
    return True


def test_full_wizard_on_complex_file(win, fixture_xlsx):
    assert run_modal(WizardDialog, _wizard_on_fixture, lambda: win.open_path(fixture_xlsx))
    wait_loaded(win)
    cols = [win.model.headerData(c, Qt.Horizontal) for c in range(win.model.columnCount())]
    assert cols == ["Layout", "Codice", "EAN", "Misto"] and win.model.total() == 60
    assert win.profile.details.code_col == "Misto" and win.profile.details.barcode_type == "ean13"

    # codice non valido per EAN-13: niente barcode, codice grande + messaggio, nessun crash
    win.search.setText("AB-3x")
    win.apply_filters()
    assert win.model.rowCount() == 1
    seen = {}
    run_modal(DetailDialog, lambda d: (seen.update(invalid=d.content.invalid, size=(d.width(), d.height())), d.accept()),
              lambda: win.show_detail(0))
    assert seen["invalid"]
    win.search.setText("COD0001")
    win.apply_filters()
    run_modal(DetailDialog, lambda d: (seen.update(invalid=d.content.invalid), d.accept()), lambda: win.show_detail(0))
    assert not seen["invalid"]


def test_second_import_skips_wizard_and_structure_change(win, fixture_xlsx, monkeypatch):
    run_modal(WizardDialog, _wizard_on_fixture, lambda: win.open_path(fixture_xlsx))
    wait_loaded(win)
    opened = []
    monkeypatch.setattr(WizardDialog, "exec", lambda self: (opened.append(1), 0)[1])
    win.open_path(fixture_xlsx)
    wait_loaded(win)
    assert not opened and win._current_toast.text().startswith("Configurazione applicata")

    from openpyxl import load_workbook
    wb = load_workbook(fixture_xlsx)
    wb["Dati"].cell(row=2, column=4, value="Posizione")
    wb.save(fixture_xlsx)
    win.open_path(fixture_xlsx)
    assert opened                                   # struttura cambiata -> wizard


def test_prefill_when_structure_changed(win, fixture_xlsx):
    run_modal(WizardDialog, _wizard_on_fixture, lambda: win.open_path(fixture_xlsx))
    wait_loaded(win)
    from openpyxl import load_workbook
    wb = load_workbook(fixture_xlsx)
    wb["Dati"].cell(row=2, column=4, value="Posizione")
    wb.save(fixture_xlsx)

    def check(d):
        d.goto(1)
        got = d.step_cols.selected_in_order()
        d.reject()
        return got

    got = run_modal(WizardDialog, check, lambda: win.open_path(fixture_xlsx))
    assert got == ["Codice", "EAN", "Misto"]        # «Layout» non esiste piu'; il resto e' precompilato
    assert win.profile.columns == ["Layout", "Codice", "EAN", "Misto"]  # annulla: stato precedente


def test_details_off_shows_info_and_opens_settings(win, fixture_xlsx):
    def basic(d):
        d.step_sheet.list.setCurrentRow(1)
        d.step_sheet.stepper.plus.click()
        d.btn_next.click()
        d.btn_next.click()                          # dettagli lasciati disattivati
        d.btn_next.click()
        return True

    run_modal(WizardDialog, basic, lambda: win.open_path(fixture_xlsx))
    wait_loaded(win)
    assert not win.profile.details.enabled
    state = {}

    def info(d):
        state["tab"] = None
        d._go()
        return True

    def settings(sd):
        state["tab"] = sd.tabs.currentIndex()
        sd.accept()

    from gui_helpers import find_dialog

    def trigger():
        orig = win.open_settings
        win.open_settings = lambda tab=0: run_modal(SettingsDialog, settings, lambda: orig(tab), delay=100)
        win.show_detail(0)

    run_modal(NoDetailsDialog, info, trigger)
    pump(300)
    assert state["tab"] == TAB_FILE


@pytest.mark.skipif(not os.path.isfile(MM16), reason="mm16.xlsx non presente")
def test_mm16_prefilled_and_second_import(win, monkeypatch):
    seen = {}

    def act(d):
        seen["step"] = d.step
        seen["cols"] = d.step_cols.selected_in_order()
        d.goto(2)
        sp = d.step_details.get_spec()
        seen["spec"] = (sp.enabled, sp.code_col, sp.barcode_type, sp.above_col, sp.hl_col, sp.hl_color)
        d.btn_next.click()
        return True

    run_modal(WizardDialog, act, lambda: win.open_path(MM16))
    wait_loaded(win)
    assert seen["step"] == 1                         # un solo foglio: salta al passo colonne
    assert seen["cols"] == ["Codice Prodotto", "Descrizione Libera", "Note Interne", "LAYOUT"]
    assert seen["spec"] == (True, "Codice Prodotto", "code128", "Descrizione Libera", "LAYOUT", "verde")
    opened = []
    monkeypatch.setattr(WizardDialog, "exec", lambda self: (opened.append(1), 0)[1])
    win.open_path(MM16)
    wait_loaded(win)
    assert not opened


def test_autoload_last_file_uses_profile(win, fixture_xlsx, monkeypatch):
    import app.main_window as mw
    run_modal(WizardDialog, _wizard_on_fixture, lambda: win.open_path(fixture_xlsx))
    wait_loaded(win)
    assert win.settings.last_file() == fixture_xlsx

    opened = []
    monkeypatch.setattr(WizardDialog, "exec", lambda self: (opened.append(1), 0)[1])
    # riavvio simulato, toggle ACCESO: carica l'ultimo file con il profilo, senza wizard
    w2 = mw.MainWindow(1.0)
    w2.resize(1024, 600)
    w2.show()
    w2.autoload()
    wait_loaded(w2)
    assert w2.model.total() == 60 and not opened
    # toggle SPENTO: stato vuoto
    w2.settings.set_remember_last_file(False)
    w3 = mw.MainWindow(1.0)
    w3.show()
    w3.autoload()
    pump(200)
    assert w3.model.total() == 0 and w3.stack.currentIndex() == 1 and not w3.search.isEnabled()
    # toggle ACCESO ma file sparito: nessun crash, stato vuoto + toast
    w2.settings.set_remember_last_file(True)
    os.remove(fixture_xlsx)
    w4 = mw.MainWindow(1.0)
    w4.show()
    w4.autoload()
    pump(200)
    assert w4.model.total() == 0 and w4._current_toast.text().startswith("File non trovato: fixture.xlsx")
    for w in (w2, w3, w4):
        w._allow_close = True
        w.close()
